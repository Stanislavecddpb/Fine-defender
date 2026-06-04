"""Генерация клиентской презентации Fine Defender (.pptx).

Запуск:
  python scripts/generate_deck.py [output.pptx]

Дизайн: фирменная палитра WB (фиолетовый/маджента), 16:9. Контент — клиентский,
на русском: проблема, решение, как работает, подключение, безопасность, цена.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --- Палитра ---
PRIMARY = RGBColor(0x4A, 0x1D, 0x6A)    # глубокий фиолетовый
ACCENT = RGBColor(0xCB, 0x11, 0xAB)     # маджента WB
DARK = RGBColor(0x21, 0x1B, 0x2E)
GRAY = RGBColor(0x6B, 0x6B, 0x76)
LIGHT = RGBColor(0xF5, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1F, 0x9D, 0x55)
RED = RGBColor(0xD6, 0x33, 0x6C)
AMBER = RGBColor(0xE8, 0xA3, 0x17)
CARD = RGBColor(0xEE, 0xE7, 0xF4)

FONT = "Segoe UI"
FONT_L = "Segoe UI Light"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# --- Хелперы ---
def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None, line_w=1.0):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.0):
    """runs: список абзацев; каждый абзац — список (txt, size, color, bold, font)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (txt, size, color, bold, font) in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = font
    return tb


def one(txt, size, color, bold=False, font=FONT):
    """Сахар: один ран в абзаце."""
    return [(txt, size, color, bold, font)]


def header(s, kicker, title, dark=False):
    """Шапка слайда: цветная плашка-акцент + надзаголовок + заголовок."""
    rect(s, 0, 0, 0.28, 7.5, ACCENT)
    base = WHITE if dark else DARK
    text(s, 0.7, 0.5, 12.0, 0.4, [one(kicker.upper(), 13, ACCENT, True)])
    text(s, 0.66, 0.85, 12.0, 1.1, [one(title, 33, base, True, FONT)])


def bullets(s, x, y, w, items, size=17, color=DARK, gap=10, marker="—  "):
    runs = []
    for it in items:
        runs.append([(marker, size, ACCENT, True, FONT), (it, size, color, False, FONT)])
    text(s, x, y, w, 5.0, runs, space_after=gap, line_spacing=1.05)


# =========================================================
# Слайд 1 — Титул
# =========================================================
s = slide(PRIMARY)
rect(s, 0, 0, 13.333, 7.5, PRIMARY)
rect(s, 0, 6.95, 13.333, 0.55, ACCENT)
rect(s, 0.9, 1.5, 1.15, 1.15, ACCENT, shape=MSO_SHAPE.OVAL)
text(s, 0.9, 1.62, 1.15, 1.0, [one("₽", 40, WHITE, True)], align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.85, 2.95, 11.5, 1.4, [one("Fine Defender", 60, WHITE, True, FONT)])
text(s, 0.9, 4.15, 11.5, 1.0,
     [one("Автоматический «Защитник от штрафов» для селлеров Wildberries", 24, RGBColor(0xE3, 0xD4, 0xEF), False, FONT_L)])
text(s, 0.9, 5.2, 11.5, 1.2,
     [one("Находим оспоримые штрафы, считаем дедлайны и готовим претензии — вы возвращаете деньги.",
          18, RGBColor(0xC9, 0xB3, 0xDC), False, FONT)])

# =========================================================
# Слайд 2 — Проблема
# =========================================================
s = slide()
header(s, "Проблема", "Маркетплейс списывает — вы узнаёте последним")
bullets(s, 0.75, 2.25, 7.4, [
    "WB удерживает штрафы прямо с баланса: повышенная логистика по обмерам, маркировка, брак при приёмке.",
    "Селлер узнаёт о списании постфактум — когда деньги уже ушли.",
    "Окно на оспаривание узкое: пропустил срок — потерял право на возврат.",
    "Доказательства приходится собирать вручную по каждому штрафу.",
], size=18, gap=16)
# Карточка-акцент справа
rect(s, 8.5, 2.3, 4.1, 3.4, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.85, 2.65, 3.5, 2.8, [
    one("Итог", 15, ACCENT, True),
    one("Значимая доля отпоримых штрафов просто теряется — процесс реактивный и ручной.", 19, PRIMARY, True),
], space_after=10, line_spacing=1.1)

# =========================================================
# Слайд 3 — Решение
# =========================================================
s = slide()
header(s, "Решение", "Fine Defender делает рутину за вас")
text(s, 0.75, 2.1, 11.8, 0.8,
     [one("Подключаете кабинет в режиме «только чтение» — сервис сам находит штрафы и готовит всё для возврата.",
          18, GRAY, False)], line_spacing=1.1)

cards = [
    ("1  Находит", "Выгружает финансовый отчёт и выделяет оспоримые штрафы среди сотен транзакций."),
    ("2  Считает", "Определяет дедлайн оспаривания и оценивает сумму к возврату."),
    ("3  Готовит", "Формирует черновик претензии и чек-лист нужных доказательств."),
    ("4  Контролирует", "Дашборд с горящими сроками и накопленной суммой «отбито ₽»."),
]
cw, gapx, x0, y0 = 2.92, 0.18, 0.75, 3.15
for i, (t, d) in enumerate(cards):
    x = x0 + i * (cw + gapx)
    rect(s, x, y0, cw, 3.1, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y0, cw, 0.12, ACCENT)
    text(s, x + 0.22, y0 + 0.3, cw - 0.44, 0.7, [one(t, 19, PRIMARY, True)])
    text(s, x + 0.22, y0 + 1.05, cw - 0.44, 2.0, [one(d, 14.5, DARK, False)], line_spacing=1.1)
text(s, 0.75, 6.55, 11.8, 0.6,
     [[("Финальную подачу делает оператор — ", 14, GRAY, False, FONT),
       ("вы ничем не рискуете.", 14, ACCENT, True, FONT)]])

# =========================================================
# Слайд 4 — Как это работает (пайплайн)
# =========================================================
s = slide(LIGHT)
header(s, "Как это работает", "Весь путь — от кабинета до готовой претензии")
steps = [
    ("Кабинет WB", "токен только\nна чтение"),
    ("Сбор отчёта", "выгрузка отчёта\nо реализации"),
    ("Классификация", "выделение\nоспоримых штрафов"),
    ("Дедлайны + черновик", "расчёт срока\nи претензия"),
    ("Дашборд", "контроль и\n«отбито ₽»"),
]
bw, aw, y = 2.05, 0.5, 3.1
total = len(steps) * bw + (len(steps) - 1) * aw
x = (13.333 - total) / 2
for i, (t, d) in enumerate(steps):
    col = PRIMARY if i % 2 == 0 else ACCENT
    rect(s, x, y, bw, 1.6, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.1, y + 0.22, bw - 0.2, 0.6, [one(t, 14.5, WHITE, True)],
         align=PP_ALIGN.CENTER)
    text(s, x + 0.1, y + 0.82, bw - 0.2, 0.7,
         [one(d.replace("\n", " "), 11, RGBColor(0xEC, 0xDD, 0xF5), False)],
         align=PP_ALIGN.CENTER, line_spacing=1.0)
    x += bw
    if i < len(steps) - 1:
        rect(s, x + 0.05, y + 0.5, aw - 0.1, 0.6, GRAY, shape=MSO_SHAPE.CHEVRON)
        x += aw
text(s, 0.75, 5.3, 11.8, 0.8,
     [[("Автоматизированы шаги 1–5. ", 15, DARK, False, FONT),
       ("Подача претензии в WB — вручную оператором (concierge-модель MVP).", 15, GRAY, False, FONT)]],
     line_spacing=1.1)

# =========================================================
# Слайд 5 — Логика классификации
# =========================================================
s = slide()
header(s, "Под капотом", "Как сервис отличает оспоримый штраф")
bullets(s, 0.75, 2.25, 7.3, [
    "Берёт транзакции, где есть удержание или штраф.",
    "По формулировке причины определяет тип штрафа.",
    "На старте — «повышенная логистика по результатам обмеров» (несоответствие реальных габаритов карточке).",
    "Для оспоримых считает дедлайн = дата начисления + окно, и оценку возврата.",
    "Остальные штрафы сохраняет для аналитики и будущих типов.",
], size=16.5, gap=13)
rect(s, 8.45, 2.3, 4.2, 3.7, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.75, 2.55, 3.65, 0.5, [one("Примеры формулировок WB", 13, ACCENT, True)])
text(s, 8.75, 3.05, 3.65, 3.0, [
    one("«Превышение габаритов по результатам обмера»", 13.5, PRIMARY, True),
    one("«Несоответствие размеров товара карточке»", 13.5, PRIMARY, True),
    one("«Габариты не соответствуют заявленным»", 13.5, PRIMARY, True),
], space_after=12, line_spacing=1.05)
text(s, 0.75, 6.5, 11.8, 0.6,
     [one("Правила и пороги вынесены в конфиг — новые типы штрафов добавляются без переписывания логики.",
          13.5, GRAY, False)])

# =========================================================
# Слайд 6 — Ключевая ценность: дедлайны (таблица)
# =========================================================
s = slide()
header(s, "Ключевая ценность", "Ни один срок не сгорит незаметно")
rows = [
    ("Товар", "Сумма", "Дедлайн", "Состояние", None),
    ("Полка настенная", "780,00 ₽", "20.05.2026", "Просрочен", RED),
    ("Коробка для хранения 50 л", "1 050,00 ₽", "05.06.2026", "Горит — 2 дня", AMBER),
    ("Органайзер для кухни", "1 480,00 ₽", "19.06.2026", "В запасе — 16 дней", GREEN),
    ("Кашпо для цветов", "920,50 ₽", "21.06.2026", "В запасе — 18 дней", GREEN),
    ("Сушилка для белья", "2 310,00 ₽", "01.07.2026", "В запасе — 28 дней", GREEN),
]
tx, ty, tw, th = 0.75, 2.3, 11.8, 3.6
gf = s.shapes.add_table(len(rows), 4, Inches(tx), Inches(ty), Inches(tw), Inches(th))
table = gf.table
table.columns[0].width = Inches(4.7)
table.columns[1].width = Inches(2.1)
table.columns[2].width = Inches(2.2)
table.columns[3].width = Inches(2.8)
for ri, row in enumerate(rows):
    table.rows[ri].height = Inches(0.6)
    for ci in range(4):
        cell = table.cell(ri, ci)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        para = cell.text_frame.paragraphs[0]
        run = para.add_run()
        run.text = row[ci]
        run.font.name = FONT
        if ri == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = PRIMARY
            run.font.color.rgb = WHITE; run.font.bold = True; run.font.size = Pt(14)
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else LIGHT
            run.font.size = Pt(13.5)
            run.font.color.rgb = DARK
            if ci == 3:
                run.font.bold = True
                run.font.color.rgb = row[4]
            if ci == 1:
                run.font.bold = True
text(s, 0.75, 6.2, 11.8, 0.7,
     [one("Сервис сортирует штрафы по близости срока и подсвечивает горящие и просроченные.",
          15, GRAY, False)])

# =========================================================
# Слайд 7 — Пример на реальном отчёте (цифры)
# =========================================================
s = slide(PRIMARY)
rect(s, 0, 0, 0.28, 7.5, ACCENT)
text(s, 0.7, 0.5, 12.0, 0.4, [one("ПРИМЕР НА ОТЧЁТЕ", 13, ACCENT, True)])
text(s, 0.66, 0.85, 12.0, 1.0, [one("Один недельный отчёт — тысячи рублей под возврат", 31, WHITE, True)])
stats = [
    ("12", "строк в отчёте"),
    ("5", "оспоримых штрафов найдено"),
    ("6 540 ₽", "потенциал возврата"),
    ("2 дня", "до ближайшего дедлайна"),
]
cw, gapx, x0, y0 = 2.92, 0.18, 0.75, 2.7
for i, (big, sub) in enumerate(stats):
    x = x0 + i * (cw + gapx)
    rect(s, x, y0, cw, 2.6, RGBColor(0x5C, 0x2A, 0x82), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y0 + 2.48, cw, 0.12, ACCENT)
    text(s, x + 0.15, y0 + 0.45, cw - 0.3, 1.1, [one(big, 38, WHITE, True)], align=PP_ALIGN.CENTER)
    text(s, x + 0.2, y0 + 1.65, cw - 0.4, 0.9, [one(sub, 14, RGBColor(0xD9, 0xC6, 0xE8), False)],
         align=PP_ALIGN.CENTER, line_spacing=1.05)
text(s, 0.75, 5.75, 11.8, 0.8,
     [one("Пример на демо-данных со структурой реального финансового отчёта Wildberries.",
          14, RGBColor(0xC2, 0xAC, 0xD6), False)])

# =========================================================
# Слайд 8 — Готовая претензия
# =========================================================
s = slide()
header(s, "Документ", "Черновик претензии — автоматически")
rect(s, 0.75, 2.25, 6.7, 4.4, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 1.05, 2.5, 6.1, 4.0, [
    one("Претензия об оспаривании штрафа за повышенную логистику", 14, PRIMARY, True),
    one("Транзакция: 200105 · Артикул: 88010005", 11.5, GRAY, False),
    one("Дата начисления: 20.04.2026 · Сумма: 780,00 ₽", 11.5, GRAY, False),
    one("Основание: «Габариты не соответствуют заявленным (обмер)»", 11.5, GRAY, False),
    one("Прошу пересмотреть результат обмера, отменить штраф и вернуть удержанные средства на баланс продавца. Фактические габариты товара соответствуют заявленным в карточке; подтверждающие материалы приложены.",
        12.5, DARK, False),
], space_after=9, line_spacing=1.12)
rect(s, 7.75, 2.25, 4.85, 4.4, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.05, 2.5, 4.25, 0.5, [one("Чек-лист доказательств", 14, ACCENT, True)])
bullets(s, 8.05, 3.05, 4.3, [
    "Фото товара с измерительной лентой",
    "Фото упаковки с лентой",
    "Упаковочный лист / спецификация",
    "Скриншот карточки с габаритами",
    "Видео процесса замера",
], size=13, gap=9, marker="✓  ")
text(s, 0.75, 6.85, 11.8, 0.5,
     [one("Текст и чек-лист настраиваются под каждый тип штрафа.", 13, GRAY, False)])

# =========================================================
# Слайд 9 — Безопасность / 152-ФЗ
# =========================================================
s = slide(LIGHT)
header(s, "Безопасность", "Ваш токен и данные под защитой")
items = [
    ("Только чтение", "Сервис запрашивает read-доступ (категория «Финансы») и ничего не меняет в вашем кабинете."),
    ("Шифрование токена", "Токен хранится в зашифрованном виде и никогда не попадает в логи."),
    ("Хостинг в РФ · 152-ФЗ", "Данные и инфраструктура размещены у российского провайдера."),
    ("Без дублей", "Повторные выгрузки идемпотентны — не плодят дубликаты и не искажают суммы."),
]
cw, gapx, gapy, x0, y0 = 5.85, 0.3, 0.3, 0.75, 2.4
for i, (t, d) in enumerate(items):
    x = x0 + (i % 2) * (cw + gapx)
    y = y0 + (i // 2) * (1.95 + gapy)
    rect(s, x, y, cw, 1.95, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y, 0.14, 1.95, ACCENT)
    text(s, x + 0.35, y + 0.25, cw - 0.6, 0.6, [one(t, 18, PRIMARY, True)])
    text(s, x + 0.35, y + 0.85, cw - 0.6, 1.0, [one(d, 14, DARK, False)], line_spacing=1.1)

# =========================================================
# Слайд 10 — Подключение
# =========================================================
s = slide()
header(s, "Подключение", "Старт за 3 шага")
steps = [
    ("Выпускаете токен", "В кабинете WB создаёте API-токен только на чтение (категория «Финансы»)."),
    ("Передаёте нам", "Заводим ваш кабинет — токен сразу шифруется и хранится защищённо."),
    ("Видите штрафы", "Сервис выгружает отчёт и показывает найденные оспоримые штрафы в дашборде."),
]
cw, gapx, x0, y0 = 3.85, 0.3, 0.75, 2.6
for i, (t, d) in enumerate(steps):
    x = x0 + i * (cw + gapx)
    rect(s, x, y0, cw, 3.0, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x + 0.3, y0 + 0.35, 0.95, 0.95, ACCENT, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.3, y0 + 0.42, 0.95, 0.8, [one(str(i + 1), 30, WHITE, True)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.3, y0 + 1.5, cw - 0.6, 0.6, [one(t, 18, PRIMARY, True)])
    text(s, x + 0.3, y0 + 2.05, cw - 0.6, 1.0, [one(d, 13.5, DARK, False)], line_spacing=1.1)
text(s, 0.75, 6.15, 11.8, 0.6,
     [[("Первые найденные штрафы — ", 16, DARK, False, FONT),
       ("в день подключения.", 16, ACCENT, True, FONT)]])

# =========================================================
# Слайд 11 — Цена
# =========================================================
s = slide(PRIMARY)
rect(s, 0, 0, 0.28, 7.5, ACCENT)
text(s, 0.7, 0.5, 12.0, 0.4, [one("МОДЕЛЬ ОПЛАТЫ", 13, ACCENT, True)])
text(s, 0.66, 0.85, 12.0, 1.0, [one("Платите только за результат", 33, WHITE, True)])
text(s, 0.75, 2.6, 11.0, 1.0,
     [one("Success fee — процент от средств, которые удалось фактически отбить.", 22, WHITE, False, FONT_L)],
     line_spacing=1.1)
pts = [
    "Нет возврата — нет оплаты.",
    "Прозрачно: в дашборде видно «под возврат» и «отбито ₽».",
    "Интересы совпадают — мы зарабатываем, только когда зарабатываете вы.",
]
runs = [[("—  ", 18, ACCENT, True, FONT), (p, 18, RGBColor(0xE3, 0xD4, 0xEF), False, FONT)] for p in pts]
text(s, 0.8, 3.9, 11.0, 2.5, runs, space_after=16, line_spacing=1.1)

# =========================================================
# Слайд 12 — Roadmap
# =========================================================
s = slide()
header(s, "Развитие", "Куда движется сервис")
items = [
    ("Сейчас", "Штрафы за повышенную логистику по результатам обмеров.", ACCENT),
    ("Далее", "Новые типы штрафов и авто-подготовка повторяющихся претензий.", PRIMARY),
    ("Перспектива", "По мере подтверждения — авто-подача и предиктивные сценарии.", GRAY),
]
y = 2.5
for t, d, col in items:
    rect(s, 0.75, y, 0.14, 1.2, col)
    text(s, 1.1, y, 3.0, 1.2, [one(t, 20, col, True)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.2, y, 8.3, 1.2, [one(d, 16.5, DARK, False)], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
    y += 1.45
text(s, 0.75, 6.7, 11.8, 0.5,
     [one("Начинаем с одного типа штрафа и расширяемся по результату — без раздувания функционала.",
          13.5, GRAY, False)])

# =========================================================
# Слайд 13 — CTA
# =========================================================
s = slide(PRIMARY)
rect(s, 0, 0, 13.333, 7.5, PRIMARY)
rect(s, 0, 0, 13.333, 0.5, ACCENT)
text(s, 0.9, 2.2, 11.5, 1.4, [one("Посмотрите свои штрафы", 46, WHITE, True)])
text(s, 0.95, 3.6, 11.0, 1.2,
     [one("Бесплатный разбор: подключим кабинет в режиме чтения и покажем найденные оспоримые штрафы по вашему ассортименту.",
          20, RGBColor(0xE3, 0xD4, 0xEF), False, FONT_L)], line_spacing=1.15)
rect(s, 0.95, 5.3, 4.6, 0.9, ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.95, 5.3, 4.6, 0.9, [one("Fine Defender", 20, WHITE, True)],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.95, 6.45, 11.0, 0.6,
     [one("Защитник от штрафов для селлеров Wildberries", 15, RGBColor(0xC2, 0xAC, 0xD6), False)])


out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Fine_Defender_презентация.pptx")
prs.save(str(out))
print(f"Сохранено: {out.resolve()}  ({len(prs.slides._sldIdLst)} слайдов)")
